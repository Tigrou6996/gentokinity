"""Prova completa senza internet: l'IA viene simulata. Uso: python test/test_offline.py"""

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config
import estrazione
import excel
import chiamate
from openpyxl import load_workbook

risposte = []
estrazione.estrai = lambda testo, giorno: risposte.pop(0)

base = {
    "chiamata_pertinente": True, "nome": "", "cognome": "", "telefono": "", "tipo_lavoro": "",
    "disponibilita": "", "competenze": "", "riassunto": "", "esito": "non_chiaro",
    "da_ricontattare": False, "data_ricontatto": "", "disponibile_dal": "",
}

with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    config.CARTELLA_BASE = tmp
    config.CARTELLA_IN_ARRIVO = tmp / "in"
    config.CARTELLA_TRASCRIZIONI = tmp / "trascrizioni"
    config.FILE_REGISTRO_ELABORATI = tmp / "gia_elaborati.txt"
    config.IGNORA_CHIAMATE_PRIMA_DEL = "2026-10-01"
    config.FILE_EXCEL = tmp / "candidati.xlsx"
    chiamate.ATTESA_FILE_STABILE = -1
    for c in (config.CARTELLA_IN_ARRIVO, config.CARTELLA_TRASCRIZIONI):
        c.mkdir()

    def chiamata(nome_file, **dati):
        (config.CARTELLA_IN_ARRIVO / nome_file).write_text("trascrizione di prova", encoding="utf-8")
        risposte.append({**base, **dati})
        chiamate.giro()

    # 1. Interessato, numero nel nome del file, cognome non detto.
    chiamata("Chiamata +39 333 1234567_20261016_101500.txt", nome="Marco", tipo_lavoro="Magazziniere",
             disponibilita="Full time, no notti", competenze="Patentino muletto",
             riassunto="Cerca lavoro come magazziniere.", esito="interessato",
             da_ricontattare=True, data_ricontatto="2026-10-22", disponibile_dal="2026-11-03")
    # 2. Non interessato: non deve finire in lista.
    chiamata("Chiamata 3471112233_20261016.txt", nome="Anna", cognome="Neri", esito="non_interessato")
    # 3. Chiamata non di lavoro: ignorata.
    chiamata("Chiamata 3200000000.txt", chiamata_pertinente=False)
    # 4. Stesso Marco richiama: la riga va aggiornata, non duplicata.
    chiamata("Chiamata 3331234567_20261020.txt", cognome="Bianchi", riassunto="Conferma l'interesse.",
             esito="interessato")
    # 5. Senza numero nel file ma detto in chiamata, da ricontattare.
    chiamata("registrazione_20261021.txt", nome="Luca", cognome="Verdi", telefono="+39 349 5556677",
             tipo_lavoro="Cameriere", riassunto="Ora lavora, sentirlo a gennaio.", da_ricontattare=True,
             data_ricontatto="2027-01-10")

    ws = load_workbook(config.FILE_EXCEL)["Candidati"]
    righe = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    for r in righe:
        print(r)

    col = {k: i for i, (_, k, _) in enumerate(excel.COLONNE)}
    assert len(righe) == 2, "attese 2 righe (Marco, Luca)"
    marco, luca = righe
    assert marco[col["nome"]] == "Marco" and marco[col["cognome"]] == "Bianchi"
    assert marco[col["telefono"]] == "+39 333 123 4567"
    assert marco[col["interessato"]] == excel.SPUNTA
    assert "[20/10/2026] Conferma" in marco[col["riassunto"]]
    assert marco[col["data_chiamata"]].date().isoformat() == "2026-10-20"
    assert marco[col["competenze"]] == "Patentino muletto"  # non cancellato dalla seconda chiamata
    assert luca[col["telefono"]] == "+39 349 555 6677"
    assert excel.formatta_telefono("+44 7700 900123") == "+447700900123"
    assert excel.formatta_telefono("0039 333 1234567") == "+39 333 123 4567"
    assert excel.formatta_telefono("02 1234567") == "+39 021234567"
    assert luca[col["disponibilita"]] == config.NON_COMUNICATO
    assert luca[col["interessato"]] == excel.VUOTA and luca[col["da_ricontattare"]] == excel.SPUNTA
    # Gli originali restano nella cartella (Google Drive) e sono segnati nel registro.
    assert len(list(config.CARTELLA_IN_ARRIVO.iterdir())) == 5
    assert len(chiamate.gia_elaborati()) == 5
    # Un nuovo giro non rielabora nulla (nessuna risposta IA in coda: se chiamasse l'IA darebbe errore).
    chiamate.giro()
    assert len(chiamate.gia_elaborati()) == 5
    # Le chiamate vecchie (prima di IGNORA_CHIAMATE_PRIMA_DEL) vengono saltate.
    (config.CARTELLA_IN_ARRIVO / "Chiamata 3331112222_20250101.txt").write_text("vecchia", encoding="utf-8")
    assert not chiamate.file_da_elaborare()
    # Nome file tipico di Xiaomi.
    assert chiamate.telefono_dal_nome_file("Mario Rossi(+393331234567)_20261016101500") == "+393331234567"
    assert chiamate.data_chiamata(Path("Mario Rossi(+393331234567)_20261016101500.mp3")).isoformat() == "2026-10-16"
    print("\nTUTTO OK")
