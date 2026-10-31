"""Scrive i dati dei candidati nel file Excel."""

import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import config

SPUNTA = "☑"
VUOTA = "☐"

FOGLIO_CANDIDATI = "Candidati"
FOGLIO_NON_INTERESSATI = "Non interessati"

# (intestazione, chiave nei dati, larghezza colonna)
COLONNE = [
    ("Data chiamata", "data_chiamata", 14),
    ("Nome", "nome", 16),
    ("Cognome", "cognome", 18),
    ("Telefono", "telefono", 16),
    ("Tipo di lavoro richiesto", "tipo_lavoro", 28),
    ("Disponibilità", "disponibilita", 32),
    ("Disponibile dal", "disponibile_dal", 15),
    ("Conoscenze / competenze", "competenze", 36),
    ("Riassunto chiamata", "riassunto", 60),
    ("Interessato", "interessato", 12),
    ("Da ricontattare", "da_ricontattare", 15),
    ("Data ricontatto", "data_ricontatto", 15),
    ("File registrazione", "file", 28),
]
INDICE = {chiave: i + 1 for i, (_, chiave, _) in enumerate(COLONNE)}
COLONNE_DATA = {"data_chiamata", "disponibile_dal", "data_ricontatto"}
COLONNE_SPUNTA = {"interessato", "da_ricontattare"}


def normalizza_telefono(numero: str) -> str:
    cifre = re.sub(r"\D", "", numero or "")
    for prefisso in ("0039", "39"):
        if cifre.startswith(prefisso) and len(cifre) - len(prefisso) >= 9:
            return cifre[len(prefisso):]
    return cifre


def formatta_telefono(numero: str) -> str:
    """Scrive tutti i numeri allo stesso modo, es. '+39 333 123 4567', così si riconoscono a colpo d'occhio."""
    numero = (numero or "").strip()
    cifre = re.sub(r"\D", "", numero)
    if not cifre:
        return ""
    estero = (numero.startswith("+") and not cifre.startswith("39")) or (
        numero.startswith("00") and not cifre.startswith("0039"))
    if estero:
        return "+" + cifre.lstrip("0")
    n = normalizza_telefono(numero)
    if len(n) == 10 and n.startswith("3"):
        return f"+39 {n[:3]} {n[3:6]} {n[6:]}"
    return f"+39 {n}"


def _a_data(valore):
    if isinstance(valore, date):
        return valore
    try:
        return date.fromisoformat(str(valore).strip())
    except ValueError:
        return None


def _crea_foglio(wb: Workbook, titolo: str):
    ws = wb.create_sheet(titolo)
    intestazione = PatternFill("solid", fgColor="1F4E78")
    for i, (testo, _, larghezza) in enumerate(COLONNE, start=1):
        cella = ws.cell(row=1, column=i, value=testo)
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = intestazione
        cella.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cella.column_letter].width = larghezza
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLONNE)).column_letter}1"

    # Menu a tendina per poter cambiare le spunte a mano, e colori quando sono spuntate.
    spunte = DataValidation(type="list", formula1=f'"{SPUNTA},{VUOTA}"', allow_blank=True)
    ws.add_data_validation(spunte)
    for chiave, colore in (("interessato", "C6EFCE"), ("da_ricontattare", "FFEB9C")):
        lettera = ws.cell(row=1, column=INDICE[chiave]).column_letter
        intervallo = f"{lettera}2:{lettera}5000"
        spunte.add(intervallo)
        ws.conditional_formatting.add(
            intervallo,
            FormulaRule(formula=[f'{lettera}2="{SPUNTA}"'], fill=PatternFill("solid", bgColor=colore)),
        )
    return ws


def apri_o_crea(percorso: Path):
    if percorso.exists():
        return load_workbook(percorso)
    wb = Workbook()
    wb.remove(wb.active)
    _crea_foglio(wb, FOGLIO_CANDIDATI)
    return wb


def _trova_riga(ws, dati: dict):
    tel = normalizza_telefono(dati["telefono"])
    nome = (dati["nome"].lower(), dati["cognome"].lower())
    for riga in range(2, ws.max_row + 1):
        valore = lambda k: str(ws.cell(row=riga, column=INDICE[k]).value or "")
        if tel and normalizza_telefono(valore("telefono")) == tel:
            return riga
        if not tel and all(nome) and (valore("nome").lower(), valore("cognome").lower()) == nome:
            return riga
    return None


def _scrivi(ws, riga: int, chiave: str, valore):
    cella = ws.cell(row=riga, column=INDICE[chiave])
    if chiave in COLONNE_DATA:
        d = _a_data(valore)
        cella.value = d if d else config.NON_COMUNICATO
        if d:
            cella.number_format = "DD/MM/YYYY"
    elif chiave in COLONNE_SPUNTA:
        cella.value = SPUNTA if valore else VUOTA
        cella.alignment = Alignment(horizontal="center", vertical="top")
        return
    else:
        cella.value = valore or config.NON_COMUNICATO
    cella.alignment = Alignment(vertical="top", wrap_text=True)


def registra(dati: dict, percorso: Path = None) -> str:
    """Aggiunge o aggiorna la riga del candidato. Restituisce cosa è stato fatto."""
    percorso = percorso or config.FILE_EXCEL
    wb = apri_o_crea(percorso)

    non_interessato = dati["esito"] == "non_interessato" and not dati["da_ricontattare"]
    ws = wb[FOGLIO_CANDIDATI]
    riga = _trova_riga(ws, dati)

    if non_interessato and riga is None:
        if not config.SALVA_NON_INTERESSATI:
            return "non interessato: non aggiunto alla lista"
        if FOGLIO_NON_INTERESSATI not in wb.sheetnames:
            _crea_foglio(wb, FOGLIO_NON_INTERESSATI)
        ws = wb[FOGLIO_NON_INTERESSATI]
        riga = _trova_riga(ws, dati)

    valori = {chiave: dati.get(chiave, "") for _, chiave, _ in COLONNE}
    valori["telefono"] = formatta_telefono(valori["telefono"])
    valori["interessato"] = dati["esito"] == "interessato"

    if riga is None:
        riga = ws.max_row + 1
        for chiave, valore in valori.items():
            _scrivi(ws, riga, chiave, valore)
        azione = "aggiunto"
    else:
        # Persona già in lista: aggiorna solo ciò che la nuova chiamata ha detto,
        # e accoda il riassunto a quelli precedenti.
        for chiave, valore in valori.items():
            precedente = ws.cell(row=riga, column=INDICE[chiave]).value
            if chiave == "riassunto" and valore and precedente and precedente != config.NON_COMUNICATO:
                giorno = _a_data(dati["data_chiamata"]).strftime("%d/%m/%Y")
                valore = f"{precedente}\n\n[{giorno}] {valore}"
            if chiave in COLONNE_SPUNTA or valore:
                _scrivi(ws, riga, chiave, valore)
        azione = "aggiornato (già presente)"
        if non_interessato:
            azione += ", ora non più interessato"

    wb.save(percorso)
    return f"{azione} nel foglio '{ws.title}'"
